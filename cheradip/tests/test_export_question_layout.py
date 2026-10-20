from io import BytesIO
import re
from unittest.mock import patch

from django.test import SimpleTestCase
from docx import Document
from docx.oxml.ns import qn
from zipfile import ZipFile
from lxml import etree
from pathlib import Path
from django.conf import settings as django_settings

from cheradip.views import ExportQuestionsView
from cheradip.views import _cq_question_structure_from_bn_markers
from cheradip.question_export_layout import apply_export_layout_snapshot, mcq_left_first_page_plan


class QuestionExportLayoutTests(SimpleTestCase):
    def test_cq_imported_marker_spacing_still_produces_four_parts(self):
        structure = _cq_question_structure_from_bn_markers('Stem। ক) A (খ) B ( গ) C ( ঘ ) D')
        self.assertEqual(structure['intro'], 'Stem।')
        self.assertEqual(structure['parts'], ['(ক) A', '(খ) B', '(গ) C', '(ঘ) D'])

    def test_cq_stem_and_four_parts_keep_preview_blocks_and_column_break(self):
        rows = [{
            'qid': f'layout-cq-{i}', 'type': 'সৃজনশীল প্রশ্ন',
            'question': text, 'answerSheetParentIndex': 0,
            'answerSheetSegmentKind': 'intro' if i == 0 else 'part',
            'answerSheetContinuation': i > 0,
            **({'answerSheetPartIndex': i - 1, 'answerSheetPartCount': 4} if i else {}),
        } for i, text in enumerate(['CQ stem', '(ক) Part A', '(খ) Part B', '(গ) Part C', '(ঘ) Part D'])]
        layout = {
            'exportLayoutVersion': 1, 'layoutColumns': 2,
            'previewQuestionsFontPxCreative': 12, 'questionsPaddingCreative': 3,
            'questionsGapCreative': 8,
            'exportPreviewPagePlan': [{'kind': 'creative', 'questionColumnIndexes': [[0, 1, 2], [3, 4]]}],
        }
        with (
            patch('cheradip.views.sync_playwright') as playwright,
            patch('cheradip.views.inject_katex_into_playwright_page'),
            patch('cheradip.views._playwright_wait_for_export_rich_images'),
        ):
            page = playwright.return_value.__enter__.return_value.chromium.launch.return_value.new_page.return_value
            page.pdf.return_value = b'%PDF-cq-segments'
            ExportQuestionsView()._build_pdf_playwright(
                rows, 'Header', 12, 12, 12, 12, 210, 297, raw_data=layout, layout_settings=layout,
            )
            html = page.set_content.call_args.args[0]
        self.assertEqual(html.count('class="q-item q-cq"'), 5)
        self.assertIn('CQ stem', html)
        self.assertNotIn('>%s</span>', html)
        self.assertIn('padding-top: 3.00px; padding-bottom: 0.00px', html)
        self.assertIn('padding-top: 0.00px; padding-bottom: 0.00px', html)
        self.assertIn('padding-top: 0.00px; padding-bottom: 3.00px', html)
        self.assertEqual(html.count('class="q-subpart-marks"'), 4)
        # The plan may break between B and C; neither exporter may rejoin them.
        result = ExportQuestionsView()._build_docx(
            rows, 'Header', 12, 12, 12, 12, 210, 297, raw_data=layout, layout_settings=layout,
        )
        document = Document(result)
        body = [p for p in document.paragraphs if 'CQ stem' in p.text or 'Part ' in p.text]
        self.assertEqual(len(body), 5)
        self.assertTrue(all(not p.paragraph_format.keep_with_next for p in body))
        texts = [p.text for p in document.paragraphs]
        b = next(i for i, text in enumerate(texts) if 'Part B' in text)
        c = next(i for i, text in enumerate(texts) if 'Part C' in text)
        self.assertTrue(any('w:type="column"' in p._p.xml for p in document.paragraphs[b:c + 1]))

    def test_docx_binding_column_keeps_ending_content_left_and_header_before_first_question_right(self):
        questions = [
            {'qid': 'first', 'type': 'সৃজনশীল প্রশ্ন', 'question': 'First question'},
            {'qid': 'last', 'type': 'সৃজনশীল প্রশ্ন', 'question': 'Ending question'},
        ]
        settings = {
            'exportLayoutVersion': 1,
            'pageSizeMcq': 'A4', 'mcqPageOrientation': 'portrait',
            'exportPreviewPagePlan': [{
                'kind': 'creative', 'leadEmpty': True, 'headerVisible': True,
                'questionColumnIndexes': [[0]], 'leadBindingIndexes': [1],
            }],
        }
        result = ExportQuestionsView()._build_docx(
            questions, 'Exam header', 12, 12, 12, 12, 210, 297,
            raw_data=settings, layout_settings=settings,
        )
        document = Document(result)
        paragraphs = document.paragraphs
        texts = [p.text for p in paragraphs]
        ending = next(i for i, text in enumerate(texts) if 'Ending question' in text)
        header = next(i for i, text in enumerate(texts) if 'Exam header' in text)
        first = next(i for i, text in enumerate(texts) if 'First question' in text)
        self.assertLess(ending, header)
        self.assertLess(header, first)
        self.assertTrue(any('w:type="column"' in p._p.xml for p in paragraphs[ending + 1:header]))

    def test_mcq_header_spans_page_before_columns_and_preserves_old_booklet_tail(self):
        plan = [
            {'kind': 'mcq', 'leadEmpty': True, 'headerInFirstColumn': True,
             'questionColumnIndexes': [[0]], 'leadBindingIndexes': [3]},
            {'kind': 'mcq', 'headerVisible': False, 'questionColumnIndexes': [[1], [2]]},
        ]
        migrated = mcq_left_first_page_plan(plan)
        self.assertTrue(plan[0]['leadEmpty'])  # never mutate saved records
        self.assertEqual([i for p in migrated for c in p['questionColumnIndexes'] for i in c], [0, 1, 2, 3])
        self.assertTrue(all(not p['headerInFirstColumn'] and not p['leadEmpty'] for p in migrated))
        questions = [{'qid': str(i), 'type': 'বহুনির্বাচনি প্রশ্ন', 'question': f'Question {i}'} for i in range(4)]
        layout = {'exportLayoutVersion': 1, 'layoutColumns': 2, 'exportPreviewPagePlan': plan}
        result = ExportQuestionsView()._build_docx(questions, 'Exam header', 12, 12, 12, 12, 210, 297,
                                                   raw_data=layout, layout_settings=layout)
        document = Document(result)
        texts = '\n'.join(p.text for p in document.paragraphs)
        self.assertLess(texts.index('Exam header'), texts.index('Question 0'))
        self.assertLess(texts.index('Question 0'), texts.index('Question 3'))
        self.assertEqual(document.sections[0]._sectPr.find(qn('w:cols')).get(qn('w:num')), '1')
        self.assertEqual(document.sections[1]._sectPr.find(qn('w:cols')).get(qn('w:num')), '2')

    def test_pdf_plan_keeps_bengali_content_and_per_kind_page_geometry(self):
        questions = [
            {
                'qid': 'cq',
                'type': 'সৃজনশীল প্রশ্ন',
                'question': 'উদ্দীপকটি পড়ে উত্তর দাও। (ক) সংজ্ঞা লেখো। (খ) ব্যাখ্যা করো।',
            },
            {
                'qid': 'mcq',
                'type': 'বহুনির্বাচনি প্রশ্ন',
                'question': 'সঠিক উত্তর কোনটি?',
                'option_1': 'প্রথম',
                'option_2': 'দ্বিতীয়',
            },
            {
                'qid': 'cq-companion',
                'type': 'গাণিতিক প্রশ্ন',
                'question': 'Companion calculation',
            },
        ]
        settings = {
            'exportLayoutVersion': 1,
            'pageSize': 'A4',
            'pageSizeCreative': 'A4',
            'pageSizeMcq': 'A5',
            'cqPageOrientation': 'portrait',
            'mcqPageOrientation': 'landscape',
            'marginTopCreative': 11,
            'marginRightCreative': 13,
            'marginBottomCreative': 17,
            'marginLeftCreative': 19,
            'marginTopMcq': 7,
            'marginRightMcq': 9,
            'marginBottomMcq': 15,
            'marginLeftMcq': 21,
            'previewQuestionsFontPxCreative': 14,
            'previewQuestionsFontPxMcq': 13,
            'previewQuestionsLineHeightCreative': 1.35,
            'previewQuestionsLineHeightMcq': 1.25,
            'questionsPaddingCreative': 7,
            'questionsPaddingMcq': 3,
            'exportPreviewPagePlan': [
                {'kind': 'creative', 'questionColumnIndexes': [[0, 2]]},
                {'kind': 'mcq', 'questionColumnIndexes': [[1]]},
            ],
        }
        with (
            patch('cheradip.views.sync_playwright') as playwright,
            patch('cheradip.views.inject_katex_into_playwright_page'),
            patch('cheradip.views._playwright_wait_for_export_rich_images'),
        ):
            page = playwright.return_value.__enter__.return_value.chromium.launch.return_value.new_page.return_value
            page.pdf.return_value = b'%PDF-layout-test'
            result = ExportQuestionsView()._build_pdf_playwright(
                questions, 'পরীক্ষা', 12, 12, 12, 12, 210, 297,
                raw_data={**settings, 'previewQuestionsFontPxCreative': 8, 'questionsPaddingCreative': 2}, layout_settings=settings,
            )
            html = page.set_content.call_args.args[0]

        self.assertEqual(result.getvalue(), b'%PDF-layout-test')
        self.assertIn('উদ্দীপকটি পড়ে উত্তর দাও', html)
        self.assertIn('সঠিক উত্তর কোনটি', html)
        self.assertIn('class="q-opt-marker">ক', html)
        self.assertNotIn('à¦', html)
        self.assertIn('size: 210.000mm 297.000mm;', html)
        self.assertIn('margin: 11.000mm 13.000mm 17.000mm 19.000mm;', html)
        self.assertIn('size: 210.000mm 148.000mm;', html)
        self.assertIn('margin: 7.000mm 9.000mm 15.000mm 21.000mm;', html)
        self.assertIn('--preview-question-lh: 1.350', html)
        self.assertIn('--preview-question-lh: 1.250', html)
        self.assertIn('--preview-q-opt-font-size: 12.00px', html)
        self.assertIn('font-size: var(--preview-q-opt-font-size, max(8px', html)
        # Extra CQ sections must not inherit the other sheet's smaller MCQ font.
        companion = re.search(r'<div class="q-item[^>]*>(?:(?!<div class="q-item).)*?Companion calculation', html, re.S)
        self.assertIsNotNone(companion)
        self.assertIn('font-size: 14.00px', companion.group(0))
        self.assertIn('--preview-question-lh: 1.350', companion.group(0))
        self.assertIn('padding-top: 7.00px; padding-bottom: 7.00px', companion.group(0))
        self.assertIn('padding-top: 3.00px; padding-bottom: 3.00px', html)
        self.assertEqual(html.count('<section class="paper '), 2)
        self.assertNotIn('page-break-after: always', html)
        self.assertIn('data:font/ttf;base64,', html)
        self.assertNotIn('file://', html)
        page.evaluate.assert_called_with('() => document.fonts.ready')

    def test_zero_selected_pdf_margin_is_not_replaced_with_default(self):
        class Request:
            data = {
                'format': 'pdf',
                'questions': [{'type': 'বহুনির্বাচনি প্রশ্ন', 'question': 'প্রশ্ন'}],
                'marginTop': 0,
                'marginRight': 0,
                'marginBottom': 0,
                'marginLeft': 0,
            }

        with patch.object(ExportQuestionsView, '_build_pdf_playwright', return_value=BytesIO(b'%PDF-test')) as render:
            response = ExportQuestionsView().post(Request())

        self.assertEqual(response.status_code, 200)
        for side in ('top', 'right', 'bottom', 'left'):
            self.assertEqual(render.call_args.kwargs['margin_%s' % side], 0)

    def test_docx_keeps_bengali_questions_and_selected_margins(self):
        questions = [{
            'qid': 'mcq',
            'type': 'বহুনির্বাচনি প্রশ্ন',
            'question': 'সঠিক উত্তর কোনটি?',
            'option_1': 'প্রথম',
            'option_2': 'দ্বিতীয়',
        }]
        settings = {
            'pageSizeMcq': 'A5',
            'previewQuestionsFontPxMcq': 8,
            'mcqPageOrientation': 'landscape',
            'marginTopMcq': 7,
            'marginRightMcq': 9,
            'marginBottomMcq': 15,
            'marginLeftMcq': 21,
            'exportPreviewPagePlan': [{'kind': 'mcq', 'questionColumnIndexes': [[0]]}],
        }
        result = ExportQuestionsView()._build_docx(
            questions, '', 12, 12, 12, 12, 210, 297,
            raw_data=settings, layout_settings=settings,
        )
        document = Document(result)
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, places=1)
        self.assertAlmostEqual(section.page_height.mm, 148, places=1)
        self.assertAlmostEqual(section.top_margin.mm, 7, places=1)
        self.assertAlmostEqual(section.right_margin.mm, 9, places=1)
        self.assertAlmostEqual(section.bottom_margin.mm, 15, places=1)
        self.assertAlmostEqual(section.left_margin.mm, 21, places=1)
        self.assertEqual(document.styles['Normal'].font.name, 'Roboto')
        self.assertIn('সঠিক উত্তর কোনটি', '\n'.join(p.text for p in document.paragraphs))
        option_runs = [run for table in document.tables for row in table.rows
                       for cell in row.cells for paragraph in cell.paragraphs for run in paragraph.runs
                       if 'প্রথম' in run.text or 'দ্বিতীয়' in run.text]
        self.assertTrue(option_runs)
        self.assertTrue(all(run.font.size.pt >= 6 for run in option_runs))

    def test_versioned_snapshot_wins_over_stale_root_without_accepting_content_fields(self):
        layout = {'exportLayoutVersion': 1, 'marginBottom': 0,
                  'previewQuestionsFontPxMcq': 10, 'questions': ['invalid'], 'format': 'bad'}
        original = {'questions': [{'question': 'Keep this'}], 'format': 'pdf',
                    'marginBottom': 25.4, 'previewQuestionsFontPxMcq': 8, 'layout_settings': layout}
        effective = apply_export_layout_snapshot(original)
        self.assertEqual(effective['marginBottom'], 0)
        self.assertEqual(effective['previewQuestionsFontPxMcq'], 10)
        self.assertEqual(effective['questions'], original['questions'])
        self.assertEqual(effective['format'], 'pdf')
        self.assertEqual(original['marginBottom'], 25.4)
        self.assertEqual(apply_export_layout_snapshot({'marginBottom': 7}, {'marginBottom': 9})['marginBottom'], 7)

    def test_docx_preserves_manual_type_size_spacing_and_embeds_matching_fonts(self):
        layout = {
            'exportLayoutVersion': 1, 'previewQuestionsFontPxMcq': 10,
            'previewQuestionsLineHeightMcq': 1.4, 'questionsPaddingMcq': 6,
            'exportPreviewPagePlan': [{'kind': 'mcq', 'headerVisible': False,
                                       'questionColumnIndexes': [[0]]}],
        }
        result = ExportQuestionsView()._build_docx(
            [{'type': 'বহুনির্বাচনি প্রশ্ন', 'question': 'Manual বাংলা text $F=ma$'}],
            '', 12, 12, 12, 12, 210, 297,
            raw_data={'previewQuestionsFontPxMcq': 8}, layout_settings=layout,
        )
        document = Document(result)
        paragraph = next(p for p in document.paragraphs if 'Manual' in p.text)
        self.assertAlmostEqual(paragraph.paragraph_format.left_indent.pt, 0)
        self.assertAlmostEqual(paragraph.paragraph_format.space_before.pt, 4.5)
        self.assertAlmostEqual(paragraph.paragraph_format.line_spacing.pt, 10.5)
        # Six vertical padding pixels remain vertical (4.5pt), including the bottom.
        self.assertAlmostEqual(document.paragraphs[-1].paragraph_format.line_spacing.pt, 4.5)
        for run in paragraph.runs:
            self.assertEqual(run.font.size.pt, 7.5)
            self.assertEqual(run._r.rPr.find(qn('w:szCs')).get(qn('w:val')), '15')
            self.assertEqual(run._r.rPr.rFonts.get(qn('w:cs')), 'Noto Sans Bengali')
        for math_run in paragraph._p.iter(qn('m:r')):
            self.assertEqual(math_run.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val')), '15')
            self.assertEqual(math_run.find(qn('w:rPr')).find(qn('w:szCs')).get(qn('w:val')), '15')
        result.seek(0)
        with ZipFile(result) as package:
            self.assertIn('w:noColumnBalance', package.read('word/settings.xml').decode())
            font_table = etree.fromstring(package.read('word/fontTable.xml'))
            for font in font_table:
                for weight, tag in [('Regular', 'embedRegular'), ('Bold', 'embedBold')]:
                    node = font.find(qn('w:' + tag))
                    if node is None:
                        continue
                    family = font.get(qn('w:name'))
                    prefix = 'NotoSansBengali' if family == 'Noto Sans Bengali' else family
                    filename = f'{prefix}-{weight}'
                    original = (Path(django_settings.BASE_DIR) / 'fonts' / (filename + '.ttf')).read_bytes()
                    embedded = bytearray(package.read('word/fonts/' + filename + '.odttf'))
                    mask = bytes.fromhex(node.get(qn('w:fontKey')).strip('{}').replace('-', ''))[::-1]
                    for index in range(32):
                        embedded[index] ^= mask[index % 16]
                    self.assertEqual(bytes(embedded), original)
            self.assertEqual(len([name for name in package.namelist() if name.endswith('.odttf')]), 4)
